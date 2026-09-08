import io
import os
import re
import hashlib
import streamlit as st
import cv2
import numpy as np
from PIL import Image, ImageDraw
from docx import Document
import torch
import torchvision
from torchvision.transforms.functional import to_tensor

st.set_page_config(page_title="AIRP", layout="centered")
st.title("AIRP")
st.markdown("Autonomous Information Redaction Privacy")

file = st.file_uploader("Upload", type=["pdf", "docx", "jpg", "jpeg", "png"])
device = "cuda" if torch.cuda.is_available() else "cpu"


def resized(image_file):
    return cv2.resize(image_file, (255, 255))


def gamma_scale(resized, gamma=1.0):
    table = np.array([((x / 255.0) ** gamma) * 255 for x in range(256)]).astype(np.uint8)
    return cv2.LUT(resized, table)


def Denoised(resized):
    return cv2.fastNlMeansDenoisingColored(resized, None, 10, 10, 7, 21)


def Imagefiles(file):
    file.seek(0)
    image_file = Image.open(file).convert("RGB")
    image_file = np.array(image_file)
    image_file = cv2.cvtColor(image_file, cv2.COLOR_RGB2BGR)
    image_file = resized(image_file)
    image_file = gamma_scale(image_file, gamma=0.8)
    image_file = Denoised(image_file)
    return Image.fromarray(cv2.cvtColor(image_file, cv2.COLOR_BGR2RGB))


@st.cache_resource
def Target_Region():
    model = torchvision.models.detection.fasterrcnn_resnet50_fpn(weights="DEFAULT")
    model.eval().to(device)
    return model


def detect_regions(image_file):
    model = Target_Region()
    with torch.no_grad():
        result = model([to_tensor(image_file).to(device)])[0]
    return [[int(v) for v in box.tolist()] for box, score in zip(result["boxes"], result["scores"]) if score.item() >= 0.60]


def Mask(image_file, box):
    image = image_file.copy()
    ImageDraw.Draw(image).rectangle(box, fill="black")
    return image


def BlackBox(image_file, box):
    return Mask(image_file, box)


def Pixelate(image_file, box):
    arr = np.array(image_file).copy(); x1, y1, x2, y2 = box; roi = arr[y1:y2, x1:x2]
    if roi.size:
        h, w = roi.shape[:2]
        small = cv2.resize(roi, (max(1, w // 12), max(1, h // 12)))
        arr[y1:y2, x1:x2] = cv2.resize(small, (w, h), interpolation=cv2.INTER_NEAREST)
    return Image.fromarray(arr)


def Blur(image_file, box):
    arr = np.array(image_file).copy(); x1, y1, x2, y2 = box; roi = arr[y1:y2, x1:x2]
    if roi.size:
        arr[y1:y2, x1:x2] = cv2.GaussianBlur(roi, (51, 51), 0)
    return Image.fromarray(arr)


def Hash(image_file, box, text="REDACTED"):
    image = image_file.copy(); value = hashlib.sha256(text.encode()).hexdigest()[:16]
    draw = ImageDraw.Draw(image); draw.rectangle(box, fill="white"); draw.text((box[0], box[1]), value, fill="black")
    return image


@st.cache_resource
def load_calamari():
    try:
        from calamari_ocr.ocr.predict.predictor import Predictor, PredictorParams
        checkpoint = "models/calamari_model"
        if not os.path.exists(checkpoint):
            return None, "Calamari checkpoint not found"
        predictor = Predictor.from_checkpoint(params=PredictorParams(), checkpoint=checkpoint)
        return predictor, "Calamari loaded"
    except Exception as e:
        return None, str(e)


def ocr_region(image_crop, predictor):
    if predictor is None:
        return ""
    try:
        arr = np.array(image_crop.convert("L"))
        preds = predictor.predict_raw([arr])
        if preds and preds[0].outputs:
            out = preds[0].outputs[0]
            return getattr(out, "sentence", getattr(out, "text", str(out)))
    except Exception:
        return ""
    return ""


@st.cache_resource
def load_clip_model():
    try:
        import clip
        model, preprocess = clip.load("ViT-B/32", device=device)
        model.eval()
        return model, preprocess
    except Exception:
        return None, None


def Choosing_Redaction(Mask, BlackBox, Pixelate, Blur, image_file, file_type="image", text=""):
    if file_type in ["docx", "pdf", "text"]:
        return "hashing" if text else "blurring"
    model, preprocess = load_clip_model()
    methods = ["blurring", "black_box", "pixelation"]
    if model is None:
        return "black_box"
    try:
        import clip
        image = preprocess(image_file).unsqueeze(0).to(device)
        prompts = clip.tokenize(["blur", "black box", "pixelation"]).to(device)
        with torch.no_grad():
            scores = model(image, prompts)[0].softmax(dim=-1)
        return methods[int(scores.argmax())]
    except Exception:
        return "black_box"


def Apply_Redaction(image_file, boxes, method, text=""):
    output = image_file.copy()
    for box in boxes:
        if method == "blurring": output = Blur(output, box)
        elif method == "black_box": output = BlackBox(output, box)
        elif method == "pixelation": output = Pixelate(output, box)
        elif method == "hashing": output = Hash(output, box, text)
        else: output = Mask(output, box)
    return output


def Veryfying_Image_Redaction(image_file):
    return len(detect_regions(image_file)) > 0


def Image_Redaction(image_file, predictor):
    output = image_file.copy()
    for _ in range(3):
        boxes = detect_regions(output)
        if not boxes:
            break
        text = ""
        for box in boxes:
            x1, y1, x2, y2 = box
            crop = output.crop((x1, y1, x2, y2))
            text += ocr_region(crop, predictor) + " "
        method = Choosing_Redaction(Mask, BlackBox, Pixelate, Blur, output, "image", text.strip())
        output = Apply_Redaction(output, boxes, method, text.strip())
        if not Veryfying_Image_Redaction(output):
            break
    return output


def Textfiles(file):
    file.seek(0)
    return "\n".join(p.text for p in Document(file).paragraphs)


def redact_text(text):
    patterns = [r"\b[\w.-]+@[\w.-]+\.\w+\b", r"\b(?:\+?\d[\d\s().-]{7,}\d)\b"]
    for pattern in patterns:
        text = re.sub(pattern, "[REDACTED]", text)
    return text


def File_redaction(file):
    file.seek(0); document = Document(file)
    text = "\n".join(p.text for p in document.paragraphs)
    method = Choosing_Redaction(Mask, BlackBox, Pixelate, Blur, Image.new("RGB", (224, 224), "white"), "docx", text)
    for paragraph in document.paragraphs:
        paragraph.text = hashlib.sha256(paragraph.text.encode()).hexdigest() if method == "hashing" else redact_text(paragraph.text)
    output = io.BytesIO(); document.save(output); output.seek(0)
    return output, method


if file is not None and st.button("Apply Redaction"):
    extension = file.name.split(".")[-1].lower()
    predictor, calamari_status = load_calamari()
    if extension in [".jpg", ".jpeg", ".png"]:
        image = Imagefiles(file); st.image(image, caption="Processed Image")
        redacted = Image_Redaction(image, predictor); st.image(redacted, caption="Redacted Image")
        out = io.BytesIO(); redacted.save(out, format="PNG")
        st.download_button("Download Redacted Image", out.getvalue(), "AIRP_redacted.png", "image/png")
    elif extension == ".docx":
        st.text_area("Text uploaded", Textfiles(file), height=200)
        out, method = File_redaction(file); st.write("Technique:", method)
        st.download_button("Download Redacted DOCX", out.getvalue(), "AIRP_redacted.docx")
    else:
        st.warning("PDF processing requires a separate PDF page-rendering pipeline.")
